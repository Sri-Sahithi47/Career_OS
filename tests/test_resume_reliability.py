"""Isolated regression tests. Never connect to the running app or its database."""

import asyncio
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient
from PyPDF2 import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from api import auth, jobs, onboarding
from api.deps import get_db
from src.database import Base
from src.models import Job, MatchedJob
from src.evaluation.tailor_service import _cache_path, TailoredResumeData
from src.recommendations import compute_resume_match
from src.settings import settings


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    monkeypatch.setattr(settings, "RESUMES_DIR", tmp_path / "resumes")
    monkeypatch.setattr(settings, "COOKIE_SECURE", False)
    monkeypatch.setattr(jobs, "refresh_user_delivery", lambda *_: None)

    def db_session():
        with sessions() as session:
            yield session

    app = FastAPI()
    for router in (auth.router, onboarding.router, jobs.router):
        app.include_router(router)
    from api.server import download_resume
    app.get("/api/download-resume/{filename}")(download_resume)
    app.dependency_overrides[get_db] = db_session
    with TestClient(app) as client:
        response = client.post('/api/v1/auth/signup', json={"username": "resume-test", "password": "test-password-123"})
        assert response.status_code == 200, response.text
        token = response.json()['token']
        user_id = response.json()['user']['id']
        client.headers['Authorization'] = f'Bearer {token}'
        with sessions() as db:
            job = Job(user_id=user_id, company="Example", title="Engineer", job_link="https://example.org/job", job_description="Python engineer", source="manual")
            db.add(job)
            db.flush()
            match = MatchedJob(user_id=user_id, job_id=job.id)
            db.add(match)
            db.commit()
            match_id = match.id
        yield client, match_id, sessions
    engine.dispose()


RESUME = "Alex Example\nalex@example.org\nEXPERIENCE\nBuilt Python services and SQL data pipelines for Example Company.\nEDUCATION\nBSc Computer Science"


def test_pasted_jobs_without_urls_are_persisted_separately(workspace):
    client, _, sessions = workspace
    ids = []
    for title in ("Frontend Engineer", "Backend Engineer"):
        response = client.post('/api/jobs', json={"title": title, "description": "Build services\nPython and SQL", "source": "manual"})
        assert response.status_code == 200, response.text
        assert response.json()['created'] is True
        assert response.json()['job']['Link'] == ''
        ids.append(response.json()['job']['id'])
    assert ids[0] != ids[1]
    with sessions() as db:
        for match_id in ids:
            match = db.query(MatchedJob).filter_by(id=match_id).one()
            assert match.job.job_description == "Build services\nPython and SQL"
            assert match.job.job_link.startswith('manual:')


@pytest.mark.parametrize('payload', [
    {"title": "Role", "description": "   "},
    {"title": " ", "description": "Description"},
    {"url": "javascript:alert(1)"},
])
def test_manual_job_validation(workspace, payload):
    client, _, _ = workspace
    assert client.post('/api/jobs', json=payload).status_code == 400


def test_extension_url_only_save_still_deduplicates(workspace):
    client, _, _ = workspace
    payload = {"url": "https://example.org/extension-role", "source": "extension"}
    first = client.post('/api/jobs', json=payload)
    second = client.post('/api/jobs', json=payload)
    assert first.status_code == second.status_code == 200
    assert first.json()['job']['id'] == second.json()['job']['id']
    assert second.json()['created'] is False


def test_upload_tailor_edit_export_round_trip(workspace):
    client, match_id, _ = workspace
    upload = client.post('/api/onboarding/resume', files={'file': ('resume.txt', RESUME, 'text/plain')})
    assert upload.status_code == 200, upload.text
    tailored = {"ats_score": 81, "location": "Remote", "tech_stack": {"Languages": ["Python"]}, "points": ["Built Python services"]}
    with patch('api.jobs.generate_tailored_resume_data', return_value=tailored) as generate:
        result = client.post(f'/api/jobs/{match_id}/tailor', json={"job_description": "Updated job description"})
    assert result.status_code == 200, result.text
    assert generate.call_args.args[0] == 'Updated job description'
    assert generate.call_args.kwargs['resume_text'] == RESUME
    data = result.json()['tailored_data']
    data['points'] = ['Reviewed & edited experience <with markup>']
    save = client.patch(f'/api/jobs/{match_id}/analysis', json=data)
    assert save.status_code == 200, save.text
    export = client.post(f'/api/jobs/{match_id}/resume')
    assert export.status_code == 200, export.text
    download = client.get(export.json()['pdf_url'])
    assert download.status_code == 200
    extracted = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(download.content)).pages)
    assert 'Alex Example' in extracted
    assert 'Reviewed & edited experience <with markup>' in extracted
    assert 'BSc Computer Science' in extracted
    assert 'Target location: Remote' in extracted
    again = client.post(f'/api/jobs/{match_id}/resume')
    assert again.status_code == 200
    assert again.json()['pdf_url'] != export.json()['pdf_url']
    assert again.json()['resume_version'] == 2
    assert len(client.get(f'/api/jobs/{match_id}/resumes').json()['resumes']) == 2


def test_resume_required_and_job_access_isolated(workspace):
    client, match_id, _ = workspace
    assert client.post(f'/api/jobs/{match_id}/tailor').status_code == 400
    assert client.post(f'/api/jobs/{match_id}/resume').status_code == 400
    response = client.post('/api/v1/auth/signup', json={"username": "different-user", "password": "test-password-123"})
    client.headers['Authorization'] = f"Bearer {response.json()['token']}"
    assert client.post(f'/api/jobs/{match_id}/tailor').status_code == 404
    assert client.get(f'/api/jobs/{match_id}/resumes').status_code == 404


def test_missing_ai_credentials_returns_service_unavailable(workspace, monkeypatch):
    client, match_id, _ = workspace
    client.post('/api/onboarding/resume', json={'resume_text': RESUME})
    monkeypatch.setattr(settings, 'GROQ_API_KEYS', [])
    response = client.post(f'/api/jobs/{match_id}/tailor')
    assert response.status_code == 503
    assert 'not configured' in response.json()['detail']


def test_download_rejects_other_users_and_anonymous_clients(workspace):
    client, match_id, _ = workspace
    assert client.post('/api/onboarding/resume', json={"resume_text": RESUME}).status_code == 200
    result = client.post(f'/api/jobs/{match_id}/resume')
    assert result.status_code == 200, result.text
    url = result.json()['pdf_url']
    response = client.post('/api/v1/auth/signup', json={"username": "other-download", "password": "test-password-123"})
    client.headers['Authorization'] = f"Bearer {response.json()['token']}"
    assert client.get(url).status_code == 404
    client.headers.pop('Authorization')
    client.cookies.clear()
    assert client.get(url).status_code == 401


def test_docx_tables_are_extracted():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = RESUME
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    text = asyncio.run(onboarding.extract_text_from_file(UploadFile(filename='resume.docx', file=buffer)))
    assert 'Built Python services' in text


@pytest.mark.parametrize('filename,content,status', [('resume.txt', b'', 400), ('resume.doc', b'legacy', 400), ('resume.pdf', b'bad PDF', 400), ('resume.txt', b'a' * (10 * 1024 * 1024 + 1), 413)])
def test_upload_errors_are_actionable(filename, content, status):
    with pytest.raises(HTTPException) as error:
        asyncio.run(onboarding.extract_text_from_file(UploadFile(filename=filename, file=BytesIO(content))))
    assert error.value.status_code == status


def test_cache_includes_candidate_context_after_800_characters(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'AI_CACHE_DIR', tmp_path)
    assert _cache_path('job', 'a' * 900 + 'one') != _cache_path('job', 'a' * 900 + 'two')


def test_malformed_ai_data_is_rejected():
    with pytest.raises(ValueError):
        TailoredResumeData(ats_score=999, location='Remote', tech_stack={'Skills': 'Python'}, points=[])


def test_partial_resume_match_does_not_saturate_to_100():
    profile = SimpleNamespace(candidate_summary='', parsed_skills=[])
    resume = SimpleNamespace(original_text='Python')
    job = SimpleNamespace(title='Engineer', search_query='', job_description='Python Rust Java', matched_skills=['Python'], missing_skills=['Rust', 'Java'])
    assert 0 < compute_resume_match(profile, resume, job)['resume_match_score'] < 100
